import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt

SRC = Path("/tmp/opencode/jesp-article/article.md")
OUT = Path("/home/sascha/Dropbox/Projects/repos/science/society/machine-consciousness/publications/scientific/en/JESP/Precaution for Machine Minds/Precaution for Machine Minds.docx")

text = SRC.read_text(encoding="utf-8")


def smart_quotes(s: str) -> str:
    s = re.sub(r"(?<=[A-Za-z])'(?=[A-Za-z])", "\u2019", s)
    out, dq_open, sq_open = [], True, True
    for ch in s:
        if ch == '"':
            out.append("\u201c" if dq_open else "\u201d")
            dq_open = not dq_open
        elif ch == "'":
            out.append("\u2018" if sq_open else "\u2019")
            sq_open = not sq_open
        else:
            out.append(ch)
    return "".join(out)


# ---- parse ----
lines = text.split("\n")
blocks, buf = [], []


def flush():
    if buf:
        blocks.append(("p", " ".join(buf)))
        buf.clear()


for raw in lines:
    line = raw.rstrip()
    if not line.strip():
        flush()
        continue
    if line.startswith("### "):
        flush()
        blocks.append(("h3", line[4:].strip()))
    elif line.startswith("## "):
        flush()
        blocks.append(("h2", line[3:].strip()))
    elif line.startswith("# "):
        flush()
        blocks.append(("h1", line[2:].strip()))
    else:
        buf.append(line.strip())
flush()

blocks = [(kind, smart_quotes(content)) for kind, content in blocks]

title = blocks[0][1]
rest = blocks[1:]

abstract_i = next(i for i, (k, c) in enumerate(rest) if k == "h2" and c == "Abstract")
keywords_i = next(i for i, (k, c) in enumerate(rest) if k == "p" and c.startswith("**Keywords"))
references_i = next(i for i, (k, c) in enumerate(rest) if k == "h2" and c == "References")

abstract_paras = [c for k, c in rest[abstract_i + 1 : keywords_i]]
keywords_line = rest[keywords_i][1]
body_blocks = rest[keywords_i + 1 : references_i]
ref_entries = [c for k, c in rest[references_i + 1 :] if k == "p"]

# ---- word count (everything except the Word-count line itself) ----
wc = sum(len(p.split()) for _, p in blocks)

# ---- build document ----
doc = Document()
cp = doc.core_properties
cp.author = ""
cp.last_modified_by = ""
cp.title = ""

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)
pf = style.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
pf.space_after = Pt(0)
pf.space_before = Pt(0)

for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)


def add_runs(par, content: str, bold=False, italic=False):
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", content):
        if not part:
            continue
        r = None
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.italic = True
            r.bold = bold
        else:
            r = par.add_run(part)
            r.bold = bold or None
            r.italic = italic or None
    return par


def para(content="", bold=False, center=False, indent=None, hang=False,
         space_before=0, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(0)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if indent is not None:
        p.paragraph_format.left_indent = Inches(indent)
    if hang:
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
    add_runs(p, content, bold=bold, italic=italic)
    return p


# ---- first page ----
para(title, bold=True, center=True, space_before=24)
para("", space_before=6)
para("Abstract", bold=True)
for c in abstract_paras:
    para(c, indent=0.5)
para("", space_before=6)
para(keywords_line)
para(f"Word count: {wc:,} (including notes and references)")

# ---- body ----
for kind, c in body_blocks:
    if kind == "h2":
        para(c, bold=True, space_before=18)
    elif kind == "h3":
        para(c, bold=True, italic=True, space_before=12)
    else:
        para(c, indent=0.5)

# ---- references ----
para("References", bold=True, space_before=18)
for c in ref_entries:
    para(c, hang=True)

# ---- footer page number ----
footer_p = doc.sections[0].footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer_p.add_run()
for el, attrs, txt in (
    ("w:fldChar", {"w:fldCharType": "begin"}, None),
    ("w:instrText", {"xml:space": "preserve"}, "PAGE"),
    ("w:fldChar", {"w:fldCharType": "end"}, None),
):
    e = OxmlElement(el)
    for a, v in attrs.items():
        e.set(qn(a), v)
    if txt:
        e.text = txt
    run._r.append(e)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"Wrote {OUT}")
print(f"Word count: {wc:,}")
